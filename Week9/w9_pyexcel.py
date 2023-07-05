import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import BarChart, Series, Reference

file_name = 'videogame/vgsales.csv'
raw_df = pd.read_csv(file_name)

#   How many games per platform
#   CREATE A FUNCTION TO DO THIS ****  
def num_games(_df, groupby_col, value_col):
    _games_df = _df.groupby([groupby_col])[value_col].count().reset_index()
    return _games_df

games_df = num_games(raw_df, "Platform", "Name")
games_df = games_df.sort_values('Name', ascending=False)[:5]
games_df.columns = ['Platform', 'NumberOfGames']
wb = Workbook()
ws = wb.active

for r in dataframe_to_rows(games_df, index=False, header=True):
    ws.append(r)


last_row = len(ws['A'])
chart1 = BarChart()
chart1.type = "col"
chart1.style = 10
chart1.title = "Video Games By Platform"
chart1.y_axis.title = '# of Games'
chart1.x_axis.title = 'Platform'

data = Reference(ws, min_col=2, min_row=1, max_row=last_row, max_col=2)
cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)
chart1.add_data(data, titles_from_data=True)
chart1.set_categories(cats)
chart1.shape = 4
ws.add_chart(chart1, "E3")


wb.save('Week9/Week9.xlsx')

# ## What country had the highest sales for N64?
# ## What were the top 5 games for the PC? 

